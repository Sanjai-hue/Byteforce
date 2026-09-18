"""Generate the sample E-Commerce SRS used for demos and evaluation.

The document deliberately contains duplicates, contradictions, ambiguities,
dependencies and missing information. The findings are NOT hard-coded anywhere in
the application — the pipeline has to discover them. The labelled expectations
used by the evaluation script live in evaluation_dataset.json.
"""

from __future__ import annotations

import json
from pathlib import Path

SAMPLES_DIR = Path(__file__).resolve().parent

TITLE = "Software Requirements Specification — ShopSphere E-Commerce Platform"

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1. Introduction",
        [
            "This document specifies the requirements for the ShopSphere e-commerce platform. "
            "It is intended for the development, quality assurance and operations teams.",
        ],
    ),
    (
        "2. Account Management",
        [
            "R001: The system shall allow new users to create an account using an email address and a password.",
            "R002: The system shall allow registered users to log in using their email address and password.",
            "R003: The system shall allow users to log in using their mobile number and password only.",
            "R004: The system shall allow a user to reset a forgotten password through a password reset link.",
            "R005: The system shall lock a user account after five consecutive failed login attempts.",
            "R006: The system shall send notifications.",
            "R007: The system shall allow users to update their profile information.",
        ],
    ),
    (
        "3. Product Catalogue",
        [
            "R008: The system shall display a catalogue of products with name, price, description and availability.",
            "R009: The system shall allow customers to search for products by keyword.",
            "R010: The system shall allow customers to filter search results by category, price range and rating.",
            "R011: Search results shall be returned quickly.",
            "R012: The system shall display product reviews submitted by verified purchasers.",
            "R013: The system shall support at least 50,000 product listings.",
        ],
    ),
    (
        "4. Shopping Cart and Checkout",
        [
            "R014: The system shall allow a logged-in customer to add products to a shopping cart.",
            "R015: The system shall allow a customer to remove products from the shopping cart.",
            "R016: The system shall calculate the cart total including applicable taxes and shipping charges.",
            "R017: The system shall allow customers to register for an account.",
            "R018: The system shall allow customers to complete a purchase using the checkout process.",
            "R019: Guest checkout is not allowed; all purchases require a registered account.",
            "R020: The system shall allow users to check out without creating an account.",
            "R021: The system shall support payment by credit card, debit card and digital wallet.",
            "R022: The system shall send an order confirmation email after a successful purchase.",
        ],
    ),
    (
        "5. Order Management",
        [
            "R023: The system shall allow customers to view their order history.",
            "R024: The system shall allow customers to track the delivery status of a placed order.",
            "R025: The system shall allow customers to cancel an order before it has been shipped.",
            "R026: The application shall respond quickly to user requests.",
            "R027: The system shall allow administrators to update the status of an order.",
            "R028: Refunds shall be processed in a reasonable time frame.",
        ],
    ),
    (
        "6. Administration",
        [
            "R029: The system shall allow administrators to add, edit and remove product listings.",
            "R030: The system shall allow administrators to view sales reports.",
            "R031: The system shall provide an appropriate level of access control for administrative functions.",
            "R032: The system shall record an audit log of all administrative actions.",
            "R033: The system shall allow administrators to manage user accounts.",
        ],
    ),
    (
        "7. Security",
        [
            "R034: The system shall encrypt all stored customer passwords.",
            "R035: All data transmitted between the client and the server shall be encrypted using TLS 1.2 or higher.",
            "R036: The system shall be secure against common web vulnerabilities.",
            "R037: The system shall require multi-factor authentication for all administrator accounts.",
            "R038: Administrators shall be able to sign in with only a username and password.",
            "R039: Customer payment card details shall not be stored on the platform's own servers.",
        ],
    ),
    (
        "8. Performance and Availability",
        [
            "R040: The system shall support at least 5,000 concurrent users without degradation of service.",
            "R041: The system should be fast during normal usage.",
            "R042: The checkout process shall complete within 3 seconds for 95% of transactions "
            "under normal operating conditions.",
            "R043: The platform shall maintain 99.9% uptime measured monthly.",
            "R044: The system shall be easy to use.",
            "R045: The system shall scale to handle large traffic volumes during sale events.",
        ],
    ),
    (
        "9. Reporting and Analytics",
        [
            "R046: The system shall generate a daily sales summary report for administrators.",
            "R047: The system shall allow administrators to export reports.",
            "R048: The system shall track product page views for analytics purposes.",
        ],
    ),
]


def build_text() -> str:
    lines = [TITLE, ""]
    for heading, items in SECTIONS:
        lines.append(heading)
        lines.append("")
        lines.extend(items)
        lines.append("")
    return "\n".join(lines)


def build_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading(TITLE, level=0)
    for heading, items in SECTIONS:
        document.add_heading(heading, level=1)
        for item in items:
            document.add_paragraph(item)
    document.save(path)


def build_pdf(path: Path) -> None:
    """Minimal single-file PDF writer, so the sample set covers all three formats."""
    lines: list[str] = [TITLE, ""]
    for heading, items in SECTIONS:
        lines.append(heading)
        for item in items:
            # Wrap so lines fit the page width.
            words = item.split()
            current = ""
            for word in words:
                if len(current) + len(word) + 1 > 95:
                    lines.append(current)
                    current = word
                else:
                    current = f"{current} {word}".strip()
            if current:
                lines.append(current)
        lines.append("")

    pages: list[list[str]] = []
    per_page = 44
    for start in range(0, len(lines), per_page):
        pages.append(lines[start : start + per_page])

    def escape(text: str) -> str:
        return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    objects: list[bytes] = []
    # Objects: 1 catalog, 2 pages tree, then (page, content) per page, then the font.
    font_id = 3 + len(pages) * 2

    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("latin-1")
    )

    for index, page_lines in enumerate(pages):
        content_id = 3 + index * 2 + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>"
            ).encode("latin-1")
        )
        body = ["BT", "/F1 9 Tf", "12 TL", "40 750 Td"]
        for line in page_lines:
            body.append(f"({escape(line)}) Tj")
            body.append("T*")
        body.append("ET")
        stream = "\n".join(body).encode("latin-1", "replace")
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )

    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, payload in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + payload + b"\nendobj\n"

    xref_at = len(out)
    count = len(objects) + 1
    out += f"xref\n0 {count}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()

    path.write_bytes(bytes(out))


# Labelled expectations for the evaluation script. These are ground-truth labels
# used only to score the pipeline; the application never reads them.
EVALUATION_DATASET = {
    "document": "ECommerce_SRS_Hackathon_Sample.docx",
    "note": (
        "Ground truth for scoring only. Requirement codes below refer to the printed "
        "labels in the sample document (R001...), which the pipeline re-numbers in "
        "document order; the evaluator maps them by matching original text."
    ),
    "duplicates": [
        {"a": "R001", "b": "R017", "reason": "account creation stated twice"},
    ],
    "contradictions": [
        {"a": "R002", "b": "R003", "reason": "email login vs mobile-number-only login"},
        {"a": "R019", "b": "R020", "reason": "guest checkout forbidden vs allowed"},
        {"a": "R037", "b": "R038", "reason": "admin MFA required vs password-only sign-in"},
    ],
    "ambiguities": ["R011", "R026", "R028", "R031", "R036", "R041", "R044", "R045"],
    "missing_information": ["R006", "R047"],
    "dependencies": [
        {"source": "R002", "target": "R001", "reason": "login requires an account"},
        {"source": "R018", "target": "R014", "reason": "checkout requires a cart"},
        {"source": "R024", "target": "R018", "reason": "tracking requires a placed order"},
    ],
    "clean_requirements": ["R042", "R043", "R035", "R013", "R040"],
}


def main() -> None:
    text = build_text()
    (SAMPLES_DIR / "ECommerce_SRS_Hackathon_Sample.txt").write_text(text, encoding="utf-8")
    build_docx(SAMPLES_DIR / "ECommerce_SRS_Hackathon_Sample.docx")
    build_pdf(SAMPLES_DIR / "ECommerce_SRS_Hackathon_Sample.pdf")
    (SAMPLES_DIR / "evaluation_dataset.json").write_text(
        json.dumps(EVALUATION_DATASET, indent=2), encoding="utf-8"
    )
    total = sum(len(items) for _, items in SECTIONS if items and items[0].startswith("R"))
    print(f"Wrote sample SRS (txt, docx, pdf) with {total} labelled requirements.")


if __name__ == "__main__":
    main()
